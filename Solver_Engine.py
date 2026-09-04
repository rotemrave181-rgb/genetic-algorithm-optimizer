import openpyxl 
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pulp import LpProblem, LpMinimize, LpVariable, lpSum, value, PULP_CBC_CMD, LpStatus

# 1. פונקציית קריאת נתונים
def load_excel_data(filepath: str):
    if not filepath.endswith('.xlsx'): #  סיומת אקסל- בדיקת בטיחות לקובץ
        filepath += '.xlsx'   
        
    wb = openpyxl.load_workbook(filepath, data_only=True) #טוענים קובץ אקסל לזכרון  
    ws = wb.active #נבחר גיליון פעיל 
    # שליפת השורות מתוך האקסל
    row_ids  = [cell.value for cell in ws[1]] # נשמור כל שורה כרשימה, רץ על כל התאים בשורה ושומר את הערך שלהם
    row_proc = [cell.value for cell in ws[2]]
    row_due  = [cell.value for cell in ws[3]]
    row_v    = [cell.value for cell in ws[4]]
    row_w    = [cell.value for cell in ws[5]]
    
    jobs = {} # יצירת מילון עם מפתח- גוב וערך- תכונות 
    job_order = [] #רשימה ששומרת את סדר הגובים ההתחלתי מהאקסל
    
    for col_idx in range(1, len(row_ids)): # עובר על כל הגובים 
        raw_id = row_ids[col_idx]
        if raw_id is None: break # תנאי עצירה- עמודה ריקה
        
        job_id = int(raw_id) # לכל עמודה שמייצגת גוב נמיר את הערכים למספרים ונארוז למילון
        jobs[job_id] = {
            'processing_time': int(row_proc[col_idx]), # מגדיר לכל גוב את הפרמטרים שלו
            'due_date': int(row_due[col_idx]),
            'earliness_penalty': int(row_v[col_idx]),
            'tardiness_penalty': int(row_w[col_idx])
        }
        job_order.append(job_id) # נוסיף את הגוב החדש לתוך הרשימה שלנו
        
    ga_params = { # - מודולרי- שומר פרמטרים של אלגוריתם גנטי 
        'pop_size': ws['A7'].value, 
        'n_elites': ws['A9'].value,
        'mut_prob': ws['A11'].value,
        'time_limit': ws['A13'].value
    }
    
    wb.close()
    return jobs, job_order, ga_params  #מחזיר את כל הנתונים שקלטנו

# 2. פונקציות הסולבר (התכנות הלינארי)

def run_optimization(jobs, job_order: list[int], time_limit_seconds: float = 5.0) -> tuple[dict, float]: # מקבלת כקלט את הגובים
    if not job_order: return {}, 0.0 # רשימה ריקה תחזיר מילון ריק
    
    
    prob = LpProblem("JIT_Scheduling", LpMinimize) #פונקציית מטרה מינימום -יוצרים מודל של תכנות לינארי
    # מגדיר משתני החלטה
    # lowBound=0- אילוץ אי שליליות
    completion = LpVariable.dicts("C", job_order, lowBound=0) 
    earliness  = LpVariable.dicts("E", job_order, lowBound=0) 
    tardiness  = LpVariable.dicts("T", job_order, lowBound=0)

    prob += lpSum( # בניית פונק המטרה שלנו
        jobs[j]['earliness_penalty'] * earliness[j] + 
        jobs[j]['tardiness_penalty'] * tardiness[j] 
        for j in job_order
    ), "Total_Penalty"

    for position, job_id in enumerate(job_order): #אילוצי הזרימה של הגובים- 
        proc_time = jobs[job_id]['processing_time'] 
        if position == 0: # עבור הגוב הראשון זמן הסיום צריך להיות גדול שווה לזמן העיבוד שלו.
            prob += completion[job_id] >= proc_time
        else: # עבור כל הגובים האחרים- זמן הסיום צריך להיות גדול שווה לזמן הסיום של הגוב הקודם ועוד זמן עיבוד של הגוב הנוכחי
            prev_id = job_order[position - 1]
            prob += completion[job_id] >= completion[prev_id] + proc_time

    for job_id in job_order: # אילוצי הקדמה ואיחור- 
        due_date = jobs[job_id]['due_date']
        prob += earliness[job_id] >= due_date - completion[job_id]
        prob += tardiness[job_id] >= completion[job_id] - due_date

    solver = PULP_CBC_CMD(msg=0, timeLimit=time_limit_seconds) # נשלח לסולבר לפתרון הבעיה
    prob.solve(solver)

    results = {} # מילון עם התוצאות
    for job_id in job_order: 
        c_j = value(completion[job_id]) # שולפת את המספר שהסולבר מצא עבור המשתנה cj
        if c_j is None:
            return _fallback_schedule(jobs, job_order)

        proc_time = jobs[job_id]['processing_time']  #שולף את כל הפרמטרים 
        due_date  = jobs[job_id]['due_date']
        s_j       = c_j - proc_time
        early_amt = max(0.0, due_date - c_j)
        tardy_amt = max(0.0, c_j - due_date)
        
        results[job_id] = { # שומרת את כל הפרמטרים במילון של התוצאות
            'Sj': round(s_j, 6), 'Cj': round(c_j, 6), 
            'Ej': round(early_amt, 6), 'Tj': round(tardy_amt, 6),
            'Pj': round(jobs[job_id]['earliness_penalty'] * early_amt + jobs[job_id]['tardiness_penalty'] * tardy_amt, 6)
        }

    total_penalty = round(sum(r['Pj'] for r in results.values()), 6) 
    # סוכמת את כל הקנסות 
    return results, total_penalty # מחזירה מילון מלא וקנס כולל

#  על פי ההנחיות אין חובה לבדוק את פלט הסולבר או לטפל בחריגת זמן.
# פונקציה זו הוספה כרשת ביטחון רק כדי למנוע קריסת תוכנית
def _fallback_schedule(jobs, job_order: list[int]) -> tuple[dict, float]: # פתרון גיבוי למקרה שהסולבר נכשל 
    results = {}
    current_time = 0.0
    for job_id in job_order:
        proc_time = jobs[job_id]['processing_time']
        due_date  = jobs[job_id]['due_date']
        c_j       = current_time + proc_time
        early_amt = max(0.0, due_date - c_j)
        tardy_amt = max(0.0, c_j - due_date)
        job_penalty = jobs[job_id]['earliness_penalty'] * early_amt + jobs[job_id]['tardiness_penalty'] * tardy_amt
        results[job_id] = {'Sj': current_time, 'Cj': c_j, 'Ej': early_amt, 'Tj': tardy_amt, 'Pj': job_penalty}
        current_time = c_j
    return results, sum(r['Pj'] for r in results.values())


# 3. פונקציות ייצוא לאקסל
def export_results(jobs, job_order, results, total_penalty, filename, original_penalty=None, num_generations='-', run_time='-'):
    if not filename.endswith('.xlsx'): filename += '.xlsx' # מייצרת קובץ אקסל 
    if original_penalty is None: original_penalty = total_penalty 

    segments = [] # מייצר את הגאנט לפי סדר הגובים
    current_time = 0.0
    for job_id in job_order:
        s_j, c_j = results[job_id]['Sj'], results[job_id]['Cj']
        if s_j > current_time + 1e-9:  # אם זמן ההתחלה של הגוב הנוכחי גדול מהזמן הנוכחי יש פער- זמן בטלה
            segments.append({'type': 'idle', 'label': f'IDLE {current_time:.1f}-{s_j:.1f}', 'start': current_time, 'end': s_j, 'job_id': None})
        segments.append({'type': 'job', 'label': f'j{job_id}', 'start': s_j, 'end': c_j, 'job_id': job_id})
        current_time = c_j

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Results'

    # עיצובים
    _IDLE_FILL = PatternFill(start_color='FFD9D9D9', end_color='FFD9D9D9', fill_type='solid')
    _BORDER = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    _BOLD = Font(bold=True, name='Arial', size=10)
    _NORM = Font(bold=False, name='Arial', size=10)

    # כותרות צד שמאלי
    row_label_map = {1: 'J', 2: 'tj', 3: 'dj', 4: 'vj', 5: 'wj', 6: 'Sj', 7: 'Cj', 8: 'Ej', 9: 'Tj', 10: 'Pj'}
    for row_idx, label in row_label_map.items():
        cell = ws.cell(row=row_idx, column=1, value=label)
        cell.font, cell.border, cell.alignment = _BOLD, _BORDER, Alignment(horizontal='center')
    ws.column_dimensions['A'].width = 14

    def _fmt(val):
        if not isinstance(val, float): return val
        return int(round(val)) if abs(round(val, 9) - round(round(val, 9))) < 1e-6 else round(val, 4)

    for seg_idx, segment in enumerate(segments):
        col = seg_idx + 2
        is_idle = (segment['type'] == 'idle')
        fill = _IDLE_FILL if is_idle else None
        j_id = segment.get('job_id')

        ws.cell(row=1, column=col, value=segment['label']).font = _BOLD
        if is_idle:
            for r in range(2, 6): ws.cell(row=r, column=col, value='-')
            ws.cell(row=6, column=col, value=_fmt(segment['start']))
            ws.cell(row=7, column=col, value=_fmt(segment['end']))
            for r in range(8, 11): ws.cell(row=r, column=col, value='-')
        else:
            ws.cell(row=2, column=col, value=jobs[j_id]['processing_time'])
            ws.cell(row=3, column=col, value=jobs[j_id]['due_date'])
            ws.cell(row=4, column=col, value=jobs[j_id]['earliness_penalty'])
            ws.cell(row=5, column=col, value=jobs[j_id]['tardiness_penalty'])
            ws.cell(row=6, column=col, value=_fmt(results[j_id]['Sj']))
            ws.cell(row=7, column=col, value=_fmt(results[j_id]['Cj']))
            ws.cell(row=8, column=col, value=_fmt(results[j_id]['Ej']))
            ws.cell(row=9, column=col, value=_fmt(results[j_id]['Tj']))
            ws.cell(row=10, column=col, value=_fmt(results[j_id]['Pj']))
        
        for r in range(1, 11):
            cell = ws.cell(row=r, column=col)
            cell.border, cell.alignment = _BORDER, Alignment(horizontal='center')
            if fill: cell.fill = fill
        ws.column_dimensions[get_column_letter(col)].width = 16

    # סיכומים
    summaries = [(12, 'Total Penalty', _fmt(total_penalty)), (13, 'Original Sol', _fmt(original_penalty)), 
                 (14, 'Total Gen Created', num_generations), (15, 'Run time', run_time)]
    for r, label, val in summaries:
        ws.cell(row=r, column=1, value=label).font = _BOLD
        ws.cell(row=r, column=2, value=val).alignment = Alignment(horizontal='center')

    # גאנט
    ws.cell(row=17, column=1, value='Gantt').font = _BOLD
    for seg_idx, segment in enumerate(segments):
        col = seg_idx + 2
        is_idle = (segment['type'] == 'idle')
        gantt_txt = f"IDLE\n({segment['start']:.1f} - {segment['end']:.1f})" if is_idle else f"j{segment['job_id']}\n({segment['start']:.1f} - {segment['end']:.1f})"
        cell = ws.cell(row=17, column=col, value=gantt_txt)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = _BORDER
        if is_idle: cell.fill = _IDLE_FILL
    ws.row_dimensions[17].height = 40

    wb.save(filename)
    print(f"[✓] Output saved to {filename}")

#  פונקציית הסולבר מודולרית- מריצה את כל תהליך 

def Run_Solver(Input_File_Name: str, Output_File_Name: str):
    print("--- Starting LP Solver Engine ---")
    try:
        jobs, job_order, ga_params = load_excel_data(Input_File_Name)
        print(f"Loaded {len(jobs)} jobs successfully.")
        
        results, total_penalty = run_optimization(jobs, job_order)
        print(f"Optimization finished. Total Penalty: {total_penalty}")
        
        export_results(jobs, job_order, results, total_penalty, Output_File_Name)
    except Exception as e:
        print(f"An error occurred in Solver Engine: {e}")